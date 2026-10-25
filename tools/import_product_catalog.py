"""Extract the supplied workbook into the versioned local development catalog.

Usage: python tools/import_product_catalog.py path/to/workbook.xlsx
Requires openpyxl. The workbook is read only; no prices or images are invented.
"""
import json
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl


def slug(value):
    value = unicodedata.normalize("NFKD", value.replace("đ", "d").replace("Đ", "D"))
    return re.sub(r"[^a-z0-9]+", "-", value.encode("ascii", "ignore").decode().lower()).strip("-")


def main():
    workbook = openpyxl.load_workbook(sys.argv[1], data_only=True, read_only=True)
    sheet = "Danh sách 100" if "Danh sách 100" in workbook.sheetnames else "Danh sách 50"
    expected = 100 if sheet.endswith("100") else 50
    rows = [row for row in workbook[sheet].iter_rows(min_row=5, values_only=True) if row[0] is not None]
    if len(rows) != expected or sorted(row[0] for row in rows) != list(range(1, expected + 1)):
        raise ValueError(f"Expected exactly {expected} unique numbered product records")
    rows.sort(key=lambda row: row[0])
    products = []
    for row in rows:
        if len(row) != 13 or any(value is None for value in row):
            raise ValueError(f"Incomplete product row: {row[0]}")
        index, category, name, brand, country, size, ingredients, description, suitable, usage, precautions, source, images = row
        products.append({
            "sku": f"CATALOG-{index:03d}",
            "name": name,
            "slug": slug(f"{brand} {name}"),
            "category": category,
            "categorySlug": slug(category),
            "brand": brand,
            "brandSlug": slug(brand),
            "brandCountry": country,
            "packageSize": size,
            "ingredients": ingredients,
            "description": description,
            "suitableFor": suitable,
            "directions": usage,
            "precautions": precautions,
            "sourceUrl": source,
            "imageSearchUrl": images,
        })
    for key in ("sku", "slug"):
        if len({product[key] for product in products}) != expected:
            raise ValueError(f"Duplicate {key}")
    for category in dict.fromkeys(row[1] for row in rows):
        detail = list(workbook[category].iter_rows(min_row=5, values_only=True))
        if [row[1:] for row in detail] != [row[1:] for row in rows if row[1] == category]:
            raise ValueError(f"Category sheet differs from master sheet: {category}")
    destination = Path(__file__).resolve().parents[1] / "src/main/resources/catalog/products.json"
    destination.parent.mkdir(parents=True, exist_ok=True)
    # Preserve separately researched prices/images when refreshing workbook text.
    enrichment_keys = ("price", "priceCurrency", "priceFrom", "priceVariant", "priceSourceUrl",
                       "priceCheckedOn", "imageUrl", "imageRemoteUrl", "imageSourceUrl", "shades", "sizes",
                       "originalPrice", "originalPriceCurrency", "conversionRate", "conversionSourceUrl")
    if destination.exists():
        existing = {product["sku"]: product for product in json.loads(destination.read_text(encoding="utf-8"))}
        for product in products:
            previous = existing.get(product["sku"], {})
            if previous.get("slug") == product["slug"]:
                product.update({key: previous[key] for key in enrichment_keys if key in previous})
    destination.write_text(json.dumps(products, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Exported {len(products)} products, {len({p['category'] for p in products})} categories, {len({p['brand'] for p in products})} brands")


if __name__ == "__main__":
    main()
